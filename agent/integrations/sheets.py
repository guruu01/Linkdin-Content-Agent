import logging
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import Dict, Any, List, Optional
from pathlib import Path
from agent.config import settings, CALENDAR_PATH, DATA_DIR, BASE_DIR
from agent.db import get_posts

logger = logging.getLogger(__name__)

HEADERS = [
    "Post ID",
    "Topic",
    "Angle / Format",
    "Hook (Above Fold)",
    "Full LinkedIn Post Content",
    "Hashtags",
    "Character Count",
    "Status",
    "Scheduled Date",
    "Created At"
]

class SheetsIntegration:
    def __init__(self):
        self.service_account_path = settings.GOOGLE_SERVICE_ACCOUNT_PATH
        self.spreadsheet_id = settings.GOOGLE_SHEETS_SPREADSHEET_ID
        self.webhook_url = settings.GOOGLE_SHEETS_WEBHOOK_URL
        self.calendar_path = CALENDAR_PATH

    @property
    def resolved_credentials_path(self) -> Optional[Path]:
        if not self.service_account_path:
            return None
        p = Path(self.service_account_path)
        if p.is_absolute() and p.exists():
            return p
        base_p = BASE_DIR / self.service_account_path
        if base_p.exists():
            return base_p
        return None

    @property
    def is_google_configured(self) -> bool:
        return bool(self.webhook_url or (self.resolved_credentials_path and self.spreadsheet_id))
        return bool(self.resolved_credentials_path and self.spreadsheet_id)

    def export_to_excel(self, posts: List[Dict[str, Any]] = None) -> str:
        """
        Generates or updates the local professional Excel content calendar.
        """
        if posts is None:
            posts = get_posts(limit=100)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LinkedIn Content Calendar"

        # Styling
        header_fill = PatternFill(start_color="0A66C2", end_color="0A66C2", fill_type="solid") # LinkedIn Blue
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )

        ws.append(HEADERS)
        for col_idx, cell in enumerate(ws[1], start=1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for p in posts:
            row = [
                p.get("id"),
                p.get("topic", ""),
                p.get("angle", ""),
                p.get("hook", ""),
                p.get("content", ""),
                p.get("hashtags", ""),
                p.get("character_count", len(p.get("content", ""))),
                p.get("status", "needs_review").replace("_", " ").title(),
                p.get("scheduled_for", ""),
                p.get("created_at", "")
            ]
            ws.append(row)

        # Style rows and column widths
        column_widths = {
            "A": 10,  # ID
            "B": 25,  # Topic
            "C": 22,  # Angle
            "D": 35,  # Hook
            "E": 65,  # Full Content
            "F": 30,  # Hashtags
            "G": 16,  # Characters
            "H": 18,  # Status
            "I": 20,  # Scheduled Date
            "J": 22   # Created At
        }

        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        for row in ws.iter_rows(min_row=2, max_col=10):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        wb.save(self.calendar_path)
        logger.info(f"Updated content calendar at {self.calendar_path}")
        return str(self.calendar_path)

    def sync_post(self, post: Dict[str, Any]) -> Dict[str, Any]:
        """
        Syncs a post draft to Google Sheets (if configured) and updates the local calendar workbook.
        """
        excel_path = self.export_to_excel()

        if self.webhook_url:
            try:
                row = [
                    post.get("id"),
                    post.get("topic", ""),
                    post.get("angle", ""),
                    post.get("hook", ""),
                    post.get("content", ""),
                    post.get("hashtags", ""),
                    len(post.get("content", "")),
                    post.get("status", "needs_review"),
                    post.get("scheduled_for", ""),
                    post.get("created_at", "")
                ]
                requests.post(self.webhook_url, json={"action": "sync_row", "row": row}, timeout=10)
                return {"success": True, "mode": "google_sheets_webhook", "excel_path": excel_path}
            except Exception as e:
                logger.warning(f"Google Sheets webhook sync failed: {e}")

        if self.resolved_credentials_path and self.spreadsheet_id:
            try:
                from google.oauth2 import service_account
                from googleapiclient.discovery import build

                scopes = ['https://www.googleapis.com/auth/spreadsheets']
                creds = service_account.Credentials.from_service_account_file(
                    str(self.resolved_credentials_path), scopes=scopes
                )
                service = build('sheets', 'v4', credentials=creds)

                values = [[
                    post.get("id"),
                    post.get("topic", ""),
                    post.get("angle", ""),
                    post.get("hook", ""),
                    post.get("content", ""),
                    post.get("hashtags", ""),
                    len(post.get("content", "")),
                    post.get("status", "needs_review"),
                    post.get("scheduled_for", "")
                ]]

                body = {'values': values}
                service.spreadsheets().values().append(
                    spreadsheetId=self.spreadsheet_id,
                    range="Sheet1!A:I",
                    valueInputOption="USER_ENTERED",
                    body=body
                ).execute()

                return {"success": True, "mode": "google_sheets_live", "excel_path": excel_path}
            except Exception as e:
                logger.warning(f"Google Sheets API call failed: {e}. Saved to Excel calendar.")

        return {
            "success": True,
            "mode": "local_excel",
            "message": "Saved to local Excel Content Calendar (data/linkedin_content_calendar.xlsx)",
            "excel_path": excel_path
        }

    def sync_batch(self, posts: List[Dict[str, Any]]) -> Dict[str, Any]:
        excel_path = self.export_to_excel(posts)

        # 1. Google Sheets via Webhook (Apps Script - No GCP key needed!)
        if self.webhook_url:
            try:
                values = []
                for post in posts:
                    values.append([
                        post.get("id"),
                        post.get("topic", ""),
                        post.get("angle", ""),
                        post.get("hook", ""),
                        post.get("content", ""),
                        post.get("hashtags", ""),
                        len(post.get("content", "")),
                        post.get("status", "needs_review"),
                        post.get("scheduled_for", ""),
                        post.get("created_at", "")
                    ])
                resp = requests.post(self.webhook_url, json={
                    "action": "sync_batch",
                    "headers": HEADERS,
                    "rows": values
                }, timeout=15)
                if resp.status_code in (200, 201, 302):
                    logger.info("Synced posts to Google Sheets via Webhook successfully!")
                    return {"success": True, "mode": "google_sheets_webhook", "count": len(posts), "excel_path": excel_path}
            except Exception as e:
                logger.warning(f"Google Sheets Webhook sync error: {e}")

        # 2. Google Sheets via Service Account
        if self.resolved_credentials_path and self.spreadsheet_id:
            try:
                from google.oauth2 import service_account
                from googleapiclient.discovery import build

                scopes = ['https://www.googleapis.com/auth/spreadsheets']
                creds = service_account.Credentials.from_service_account_file(
                    str(self.resolved_credentials_path), scopes=scopes
                )
                service = build('sheets', 'v4', credentials=creds)

                # Check if headers exist
                sheet_data = service.spreadsheets().values().get(
                    spreadsheetId=self.spreadsheet_id, range="Sheet1!A1:J1"
                ).execute()
                has_headers = bool(sheet_data.get('values'))

                values = []
                if not has_headers:
                    values.append(HEADERS)

                for post in posts:
                    values.append([
                        post.get("id"),
                        post.get("topic", ""),
                        post.get("angle", ""),
                        post.get("hook", ""),
                        post.get("content", ""),
                        post.get("hashtags", ""),
                        len(post.get("content", "")),
                        post.get("status", "needs_review"),
                        post.get("scheduled_for", ""),
                        post.get("created_at", "")
                    ])

                if values:
                    service.spreadsheets().values().append(
                        spreadsheetId=self.spreadsheet_id,
                        range="Sheet1!A:J",
                        valueInputOption="USER_ENTERED",
                        body={'values': values}
                    ).execute()

                logger.info(f"Successfully synced {len(posts)} posts to Google Sheet {self.spreadsheet_id}")
                return {"success": True, "mode": "google_sheets_live", "count": len(posts), "excel_path": excel_path}
            except Exception as e:
                logger.warning(f"Google Sheets batch sync failed: {e}. Saved to local Excel calendar.")

        return {
            "success": True,
            "mode": "local_excel",
            "count": len(posts),
            "excel_path": excel_path
        }

sheets_client = SheetsIntegration()

